import ctypes
import os
import re
import shutil
import threading
import time
from ctypes import wintypes
from pathlib import Path

import openpyxl
import pandas as pd
from camoufox.sync_api import Camoufox
from playwright.sync_api import expect

TARGET_URL = "https://fasih-sm.bps.go.id/app"
# Each script reads its own folder: kbli.py -> input\kbli\
INPUT_DIR = Path(__file__).parent / "input" / Path(__file__).stem
ASSIGNMENT_URL = "https://fasih-sm.bps.go.id/app/assignment/fd68e454-ba45-4b85-8205-f3bf777ded24/"
ID_COLUMN = "assignment_id"
HEADER_ROW = 2  # Excel row with the column names (row 1 holds group headers)
# Result columns added at the end of the sheet: HASIL is "sukses"/"gagal" (empty = not run yet)
HASIL_COLUMN = "hasil"
PESAN_COLUMN = "pesan"
SUKSES = "sukses"
GAGAL = "gagal"
VALIDASI_MSG = "kena error validasi, silakan cek kembali isiannya"
PAGE_LOAD_TIMEOUT = 60000  # ms
KBLI_TYPE_DELAY = 2000  # ms to wait after typing a KBLI code before selecting
KBLI_HOVER_DELAY = 1000  # ms to wait after hovering the KBLI option, before clicking
KBLI_CLICK_DELAY = 300  # ms to wait after clicking the KBLI option, before going on to the remarks

# Assignment page selectors taken from otomasi-force-submit
QUESTIONER_SELECTOR = "xpath=/html/body/div[2]/div/div[1]/div/div/div[1]/div[2]/div[2]/div[2]/div[1]/div/div/strong"
EDIT_BTN_SELECTOR = "xpath=/html/body/div[1]/div/div/div/div[2]/div/div/div/button[5]"
EDIT_MODE_SELECTOR = "xpath=/html/body/div[1]/div/div/div/div[2]/div/div[2]/div/p"
# The old absolute xpath landed on "SE2026 - P"; match the sidebar item by title instead
BUSINESS_SECTION_SELECTOR = "xpath=//aside//div[contains(@class,'tw:cursor-pointer') and @title='SE2026 - L BLOK II']"
CATATAN_SECTION_SELECTOR = "xpath=//aside//div[contains(@class,'tw:cursor-pointer') and @title='CATATAN']"
# "Tampilkan Anomali Usaha dan Keluarga" switch in the CATATAN section
ANOMALI_SWITCH_SELECTOR = "#cek_anomali_button input[role=switch]"
# Each usaha listed in the sidebar under "Keterangan Usaha/Perusahaan"
USAHA_ITEM_SELECTOR = "xpath=//div[@title='Keterangan Usaha/Perusahaan']/following-sibling::div//div[contains(@class,'tw:cursor-pointer') and @title]"
# First input after the question 13 label = 13a (kegiatan utama)
Q13_INPUT_SELECTOR = "xpath=//b[starts-with(normalize-space(.),'13. Kegiatan utama')]/following::input[1]"
KEG_UTAMA_COLUMN = "kbli_keg_utama"

# Excel question code -> FASIH question variable. Each question is rendered as a
# container with id "<variable>#<row>", e.g. "layanan_mamin#1002".
QUESTION_VARS = {
    "9a": "jenis_usaha",
    "13b1": "produk_sendiri",
    "13b2": "layanan_mamin",
    "13b3": "keg_penjualan",
    "13b4jasa": "keg_jasa",
    "13b4pertanian": "keg_jasa",
    "kat_kbli": "kbli",
}
# PERBAIKAN columns, in form order (13b3/13b4 only appear after the earlier answers)
PERBAIKAN_COLUMNS = ["13b1", "13b2", "13b3", "13b4jasa", "13b4pertanian", "kat_kbli"]
# CATATAN headers repeat the PERBAIKAN names, so pandas suffixes them with ".1"
CATATAN_COLUMNS = {
    "9a": "9a", "13b1": "13b1.1", "13b2": "13b2.1", "13b3": "13b3.1",
    "13b4jasa": "13b4jasa.1", "13b4pertanian": "13b4pertanian.1", "kat_kbli": "kat_kbli.1",
}
# 13b4 is a single question: 1 = Jasa, 2 = Pertanian, Perikanan, dan Kehutanan
Q13B4_VALUES = {"13b4jasa": "1", "13b4pertanian": "2"}

# The user logs in in the browser; the script waits until FASIH shows the survey list
LOGGED_IN_URL = re.compile(r"^https://fasih-sm\.bps\.go\.id/app/surveys")
LOGIN_TIMEOUT = 15 * 60 * 1000  # ms the user gets to finish logging in


# A browser window that is behind other windows gets throttled by Firefox; these keep it
# running at full speed
BROWSER_PREFS = {
    "widget.windows.window_occlusion_tracking.enabled": False,
    "dom.min_background_timeout_value": 0,
    "dom.min_background_timeout_value_without_budget_throttling": 0,
    "dom.timeout.enable_budget_timer_throttling": False,
}


def _is_camoufox_window(hwnd):
    cls = ctypes.create_unicode_buffer(64)
    ctypes.windll.user32.GetClassNameW(hwnd, cls, 64)
    if cls.value != "MozillaWindowClass":
        return False
    pid = wintypes.DWORD()
    ctypes.windll.user32.GetWindowThreadProcessId(hwnd, ctypes.byref(pid))
    proc = ctypes.windll.kernel32.OpenProcess(0x1000, False, pid.value)  # PROCESS_QUERY_LIMITED_INFORMATION
    if not proc:
        return False
    try:
        buf = ctypes.create_unicode_buffer(260)
        size = wintypes.DWORD(260)
        ctypes.windll.kernel32.QueryFullProcessImageNameW(proc, 0, buf, ctypes.byref(size))
        return os.path.basename(buf.value).lower() == "camoufox.exe"
    finally:
        ctypes.windll.kernel32.CloseHandle(proc)


def keep_browser_unminimized():
    """A minimized window stops painting, so every click hangs. Un-minimize Camoufox
    (without taking focus; it stays behind other windows) whenever that happens."""
    user32 = ctypes.windll.user32
    callback_type = ctypes.WINFUNCTYPE(wintypes.BOOL, wintypes.HWND, wintypes.LPARAM)

    def restore(hwnd, _):
        if user32.IsIconic(hwnd) and _is_camoufox_window(hwnd):
            user32.ShowWindow(hwnd, 4)  # SW_SHOWNOACTIVATE
        return True

    callback = callback_type(restore)

    def loop():
        while True:
            user32.EnumWindows(callback, 0)
            time.sleep(1)

    threading.Thread(target=loop, daemon=True).start()


def login(page):
    """The user logs in to FASIH in the browser themselves; wait until that is done."""
    page.goto(TARGET_URL, timeout=60000)
    print()
    print("=" * 70)
    print(" Silakan LOGIN di jendela browser yang terbuka")
    print(" (captcha jika ada, Lanjutkan dengan SSO, username, password, OTP).")
    print(" Script akan lanjut otomatis setelah login berhasil.")
    print("=" * 70)
    # Logged out, FASIH stays on /app ("Selamat Datang Kembali"); after login it
    # lands on the survey list at /app/surveys
    try:
        page.wait_for_url(LOGGED_IN_URL, timeout=LOGIN_TIMEOUT)
    except Exception:
        raise SystemExit("[!] Login belum selesai dalam 15 menit (atau browser ditutup). Jalankan lagi.")
    print("[+] Login berhasil")


class RowFailed(Exception):
    """A row that could not be sent; the message goes into the PESAN column."""


class ResultSheet:
    """Writes HASIL / PESAN for each row back into the input Excel file."""

    def __init__(self, path):
        self.path = path
        # Keep a copy of the file as it was before this run, in input\<script>\backup\
        backup_dir = path.parent / "backup"
        backup_dir.mkdir(exist_ok=True)
        shutil.copy2(path, backup_dir / f"{path.stem} {time.strftime('%Y%m%d-%H%M%S')}{path.suffix}")
        self.wb = openpyxl.load_workbook(path)
        self.ws = self.wb.active
        headers = [c.value for c in self.ws[HEADER_ROW]]
        # Add the two columns at the end the first time the script runs on this file
        self.cols = {}
        for name in (HASIL_COLUMN, PESAN_COLUMN):
            if name not in headers:
                headers.append(name)
                self.ws.cell(row=HEADER_ROW, column=len(headers), value=name)
            self.cols[name] = headers.index(name) + 1

    def mark(self, excel_row, hasil, pesan=""):
        self.ws.cell(row=excel_row, column=self.cols[HASIL_COLUMN], value=hasil)
        self.ws.cell(row=excel_row, column=self.cols[PESAN_COLUMN], value=pesan)
        # Save after every row so progress survives a crash. Write a temporary file first
        # and swap it in only when complete, so a failed save never damages the Excel file.
        tmp = self.path.with_name(self.path.name + ".tmp")
        try:
            self.wb.save(tmp)
            while True:
                try:
                    os.replace(tmp, self.path)
                    return
                except PermissionError:
                    input(f"[!] {self.path.name} sedang dibuka di Excel. Tutup file itu, lalu tekan Enter...")
        finally:
            tmp.unlink(missing_ok=True)


def read_rows():
    files = sorted(f for f in INPUT_DIR.glob("*.xlsx") if not f.name.startswith("~$"))
    if not files:
        raise SystemExit(f"Tidak ada file .xlsx di {INPUT_DIR}")
    # Row 1 holds group headers; actual column names are on row 2
    # dtype=str keeps leading zeros in KBLI codes (e.g. "01309")
    df = pd.read_excel(files[0], header=HEADER_ROW - 1, dtype=str).dropna(subset=[ID_COLUMN])
    print(f"[+] Membaca {files[0].name}")
    # df index 0 is the first row under the header, i.e. Excel row HEADER_ROW + 1
    rows = [(i + HEADER_ROW + 1, ASSIGNMENT_URL + row[ID_COLUMN].strip(), row) for i, row in df.iterrows()]
    return files[0], rows


def cell(row, column):
    """Stripped cell value, or None when the cell is empty."""
    value = row.get(column)
    if pd.isna(value) or not str(value).strip():
        return None
    return str(value).strip()


def normalize(text):
    return " ".join(str(text).split()).casefold()


def open_matching_usaha(page, keg_utama, tag):
    """Click each usaha until one whose 13a equals keg_utama; leave that one open."""
    items = page.locator(USAHA_ITEM_SELECTOR)
    items.first.wait_for(timeout=PAGE_LOAD_TIMEOUT)
    found = []
    for j in range(items.count()):
        item = items.nth(j)
        name = item.get_attribute("title")
        item.click()
        # The clicked usaha gets the active (bg-primary) class once its form is shown
        page.wait_for_function(
            "el => el.className.includes('tw:bg-primary')",
            arg=item.element_handle(),
            timeout=PAGE_LOAD_TIMEOUT,
        )
        q13 = page.locator(Q13_INPUT_SELECTOR)
        q13.wait_for(timeout=PAGE_LOAD_TIMEOUT)
        value = q13.input_value()
        found.append(f"{name}: {value}")
        if normalize(value) == normalize(keg_utama):
            print(f"{tag} Usaha cocok: {name}")
            return True
    print(f"{tag} Tidak ada usaha dengan 13a = '{keg_utama}'. Ditemukan: {found}")
    return False


def question(page, code):
    return page.locator(f'[id^="{QUESTION_VARS[code]}#"]')


def select_radio(page, code, value, tag):
    q = question(page, code)
    q.wait_for(timeout=10000)  # 13b3/13b4 only appear after the skip rule is met
    radio = q.locator(f'input[type=radio][value="{value}"]')
    if radio.is_checked():
        print(f"{tag} {code} sudah = {value}")
        return
    # The radio input itself ignores clicks; its <label for=...> is what FASIH listens to
    page.click(f'label[for="{radio.get_attribute("id")}"]')
    expect(radio).to_be_checked()
    print(f"{tag} {code} = {value}")


def select_kbli(page, code_kbli, tag):
    box = question(page, "kat_kbli").locator("textarea")
    box.wait_for(timeout=10000)
    # The selected KBLI is shown as the placeholder, e.g. "[G][47249]Perdagangan ..."
    selected = lambda: f"[{code_kbli}]" in (box.get_attribute("placeholder") or "")
    if selected():
        print(f"{tag} kat_kbli sudah = {code_kbli}")
        return
    # Each option carries its code in data-key
    option = page.locator(f'[role=listbox] [role=option][data-key="{code_kbli}"]')
    # FASIH selects on pointerdown, and pressing an already-selected option clears the
    # field ("Pilih salah satu..."). So never press a selected option, and retry if the
    # field ended up empty.
    for attempt in range(1, 4):
        box.click()
        box.fill("")  # on a retry the box may still hold the previous search text
        # Type key by key so the combobox filters exactly like manual typing
        box.press_sequentially(code_kbli, delay=100)
        # Give the list time to filter on the typed code before picking an option
        page.wait_for_timeout(KBLI_TYPE_DELAY)
        try:
            option.wait_for(timeout=10000)
        except Exception:
            shown = page.locator("[role=listbox] [role=option]").evaluate_all(
                "els => els.map(e => e.getAttribute('data-key'))"
            )
            raise RuntimeError(f"Opsi KBLI {code_kbli} tidak muncul (diketik: '{box.input_value()}', opsi: {shown})")
        # Hover first, pause, then click, pause, so the list is settled at every step
        option.hover()
        page.wait_for_timeout(KBLI_HOVER_DELAY)
        if option.get_attribute("aria-selected") != "true":
            option.click()
            page.wait_for_timeout(KBLI_CLICK_DELAY)
        try:
            expect(box).to_have_attribute("placeholder", re.compile(re.escape(f"[{code_kbli}]")), timeout=3000)
            break
        except AssertionError:
            print(f"{tag} kat_kbli percobaan {attempt} gagal (terisi: '{box.get_attribute('placeholder')}'), ulangi...")
    if not selected():
        raise RuntimeError(f"kat_kbli gagal diisi {code_kbli}")
    # The dropdown sometimes stays open after selecting; it is also a [role=dialog]
    # and would get in the way of the Catatan dialog. Escape closes it and keeps the value.
    listbox = page.locator("[role=listbox]")
    if listbox.count():
        page.keyboard.press("Escape")
        try:
            listbox.wait_for(state="detached", timeout=2000)
        except Exception:
            pass  # still closing; the Catatan step picks its own dialog anyway
    print(f"{tag} kat_kbli = {code_kbli}")


def add_catatan(page, code, text, tag):
    q = question(page, code)
    # The answers are already filled, so a question that isn't shown by now won't appear
    if not q.count():
        raise RuntimeError("pertanyaan tidak tampil (skip rule)")
    # Each question has two Catatan buttons (one per layout); only one is visible
    q.locator("button[aria-label=Catatan]:visible").first.click()
    dialog = page.locator("[role=dialog]", has=page.locator("textarea[placeholder^='Tambah catatan']"))
    box = dialog.locator("textarea")
    box.fill(text)
    dialog.locator("button[title=Simpan]").click()
    # The textarea is cleared once the note is posted
    page.wait_for_function("el => el.value === ''", arg=box.element_handle(), timeout=10000)
    dialog.locator("button[aria-label=Dismiss]").click()
    dialog.wait_for(state="detached", timeout=10000)
    print(f"{tag} Catatan {code} = {text}")


def apply_perbaikan(page, row, tag):
    for code in PERBAIKAN_COLUMNS:
        value = cell(row, code)
        if value is None:
            continue
        if code == "kat_kbli":
            select_kbli(page, value.zfill(5), tag)
        elif code in Q13B4_VALUES:
            select_radio(page, code, Q13B4_VALUES[code], tag)
        else:
            select_radio(page, code, value, tag)


def apply_catatan(page, row, tag):
    """Write every CATATAN remark; returns the codes of the ones that failed."""
    failed = []
    for code, column in CATATAN_COLUMNS.items():
        text = cell(row, column)
        if text is None:
            continue
        # A question hidden by the skip rule (e.g. 13b4 when 13b1 = Ya) can't take a
        # remark; skip it without blocking Kirim
        if not question(page, code).count():
            print(f"{tag} Catatan {code} dilewati: pertanyaan tidak tampil (skip rule)")
            continue
        # One remark failing must not stop the rest
        try:
            add_catatan(page, code, text, tag)
        except Exception as e:
            print(f"{tag} Catatan {code} gagal: {e}")
            failed.append(code)
    return failed


def enable_anomali(page, tag):
    page.click(CATATAN_SECTION_SELECTOR)
    switch = page.locator(ANOMALI_SWITCH_SELECTOR)
    switch.wait_for(state="attached", timeout=PAGE_LOAD_TIMEOUT)
    if not switch.is_checked():
        # The real input is 1x1 px; the visible toggle is the div right after it
        page.click(ANOMALI_SWITCH_SELECTOR + " + div")
        expect(switch).to_be_checked()
    print(f"{tag} Tampilkan Anomali Usaha dan Keluarga = aktif")


def kirim(page, tag):
    """Open the Kirim summary and send only when it reports 0 GALAT; raises RowFailed otherwise."""
    page.locator("button", has_text=re.compile(r"^\s*Kirim\s*$")).first.click()
    dialog = page.locator("[role=dialog]", has_text="Jawaban")
    dialog.wait_for(timeout=PAGE_LOAD_TIMEOUT)
    # The summary has cards like "GALAT 2 Perlu diperbaiki"; GALAT blocks submission
    galat_text = dialog.locator("button", has_text="GALAT").inner_text()
    galat = int(re.search(r"\d+", galat_text).group())
    if galat:
        print(f"{tag} Tidak dikirim: {galat} GALAT harus diperbaiki dulu.")
        dialog.locator("button", has_text=re.compile(r"^\s*Batal\s*$")).click()
        dialog.wait_for(state="detached", timeout=10000)
        raise RowFailed(VALIDASI_MSG)
    dialog.locator("button", has_text=re.compile(r"^\s*Kirim\s*$")).click()
    # Second step: "Konfirmasi Kirim - Apakah Anda yakin ingin mengirimkan data ini?"
    confirm = page.locator("[role=dialog]", has_text="Konfirmasi Kirim")
    confirm.wait_for(timeout=PAGE_LOAD_TIMEOUT)
    confirm.locator("button", has_text=re.compile(r"^\s*Konfirmasi\s*$")).click()
    # Success shows "Assignment Berhasil Disubmit"; the server can take a while
    page.locator("[role=alertdialog]", has_text="Berhasil Disubmit").wait_for(timeout=PAGE_LOAD_TIMEOUT)
    print(f"{tag} Terkirim.")


def process_assignment(page, url, row, tag):
    page.goto(url, wait_until="domcontentloaded", timeout=PAGE_LOAD_TIMEOUT)

    # Step 1: wait for the questioner name to load (page is ready)
    page.wait_for_selector(QUESTIONER_SELECTOR, timeout=PAGE_LOAD_TIMEOUT)
    print(f"{tag} Questioner loaded.")

    # Step 2: click the edit button (5th button in that toolbar)
    page.click(EDIT_BTN_SELECTOR)
    print(f"{tag} Clicked edit button.")

    # Step 3: wait for the page to reload into edit mode
    # (this element only appears after the reload finishes)
    page.wait_for_selector(EDIT_MODE_SELECTOR, timeout=PAGE_LOAD_TIMEOUT)
    print(f"{tag} Now in edit mode.")

    # Step 4: click the business section (in the sidebar)
    page.click(BUSINESS_SECTION_SELECTOR)
    print(f"{tag} Clicked business section.")

    # Step 5: open the usaha whose question 13a matches kbli_keg_utama
    keg_utama = cell(row, KEG_UTAMA_COLUMN)
    if not open_matching_usaha(page, keg_utama, tag):
        raise RowFailed(f"tidak ada usaha dengan 13a = '{keg_utama}'")

    # Step 6: fill the PERBAIKAN answers (raises on failure), then write the CATATAN remarks
    apply_perbaikan(page, row, tag)
    failed_catatan = apply_catatan(page, row, tag)

    # Step 7: in the CATATAN section, turn on "Tampilkan Anomali Usaha dan Keluarga"
    enable_anomali(page, tag)
    if failed_catatan:
        raise RowFailed(f"catatan gagal ditulis: {', '.join(failed_catatan)}")

    # Step 8: send, but only when the Kirim summary has no GALAT
    kirim(page, tag)


def main():
    keep_browser_unminimized()
    with Camoufox(headless=False, humanize=True, firefox_user_prefs=BROWSER_PREFS) as browser:
        page = browser.new_page()

        login(page)

        path, rows = read_rows()
        sheet = ResultSheet(path)
        # Only rows that failed or haven't run yet; "sukses" rows are already sent
        todo = [r for r in rows if cell(r[2], HASIL_COLUMN) != SUKSES]
        print(f"[+] {len(rows) - len(todo)} baris sudah sukses, {len(todo)} baris akan diproses")
        for i, (excel_row, url, row) in enumerate(todo, 1):
            tag = f"[{i}/{len(todo)}]"
            print(f"{tag} {url}")
            try:
                process_assignment(page, url, row, tag)
                sheet.mark(excel_row, SUKSES)
            except Exception as e:
                # Browser closed by the user: stop here and leave this row and the rest
                # empty (not run yet), instead of marking them all "gagal"
                if page.is_closed() or not browser.is_connected():
                    print(f"{tag} Browser ditutup. Proses dihentikan; baris ini dan sisanya tetap kosong.")
                    break
                # RowFailed carries a ready message; anything else keeps its first line
                pesan = str(e) if isinstance(e, RowFailed) else (str(e).strip().splitlines() or [type(e).__name__])[0]
                print(f"{tag} Gagal: {pesan}")
                sheet.mark(excel_row, GAGAL, pesan)

        input("Selesai. Tekan Enter untuk menutup browser...")


if __name__ == "__main__":
    main()
